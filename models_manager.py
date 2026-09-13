# models_manager.py
"""
Módulo para la gestión del catálogo de modelos y la generación del reporte
Excel de ventas desglosado por modelo, color y tallas.
"""

import json
import re
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from models import DayBatch, OrderItem

CATALOG_PATH = Path(__file__).resolve().parent / "models_catalog.json"

DEFAULT_COLORS: Dict[str, str] = {
    "01": "NEGRO",
    "02": "ROJO",
    "03": "BLANCO",
    "04": "MARINO",
    "05": "AZUL REY",
    "06": "OXFORD",
    "07": "GRIS",
    "07O": "GRIS OSCURO",
    "08": "BEIGE",
    "09": "NARANJA",
    "10": "AMARILLO",
    "11": "CELESTE",
    "12": "AZUL",
    "13": "TURQUESA",
    "14": "CORAL",
    "15": "VINO",
    "16": "PALO DE ROSA",
    "17": "VERDE",
    "20": "PETROLEO",
    "21": "MANZANA",
}

SIZE_NORMALIZATION: Dict[str, str] = {
    "S": "CH",
    "CH": "CH",
    "CHICA": "CH",
    "SMALL": "CH",
    "M": "M",
    "MED": "M",
    "MEDIANA": "M",
    "MEDIUM": "M",
    "L": "G",
    "G": "G",
    "GDE": "G",
    "GRANDE": "G",
    "LARGE": "G",
    "XL": "XG",
    "XG": "XG",
    "EXTRA GRANDE": "XG",
    "XXL": "2XG",
    "2XL": "2XG",
    "2XG": "2XG",
    "UNITALLA": "Unitalla",
    "UNI": "Unitalla",
    "U": "Unitalla",
}

DEFAULT_CATALOG = {
    "BARCELONA": {
        "aliases": ["BARCELONA", "BCN", "COPO26BCN"],
        "colors": [
            {"name": "NEGRO", "code": "01"},
            {"name": "MARINO", "code": "04"},
            {"name": "BLANCO", "code": "03"},
            {"name": "VINO", "code": "15"},
        ],
        "sizes": ["CH", "M", "G", "XG"],
    },
    "JAPON": {
        "aliases": ["JAPON", "JPN", "COPO26JPN"],
        "colors": [
            {"name": "NEGRO", "code": "01"},
            {"name": "MARINO", "code": "04"},
            {"name": "BLANCO", "code": "03"},
            {"name": "ROJO", "code": "02"},
        ],
        "sizes": ["CH", "M", "G", "XG"],
    },
    "MIAMI": {
        "aliases": ["MIAMI", "MIA", "COPO26MIA"],
        "colors": [
            {"name": "NEGRO", "code": "01"},
            {"name": "BLANCO", "code": "03"},
            {"name": "MARINO", "code": "04"},
            {"name": "VERDE", "code": "17"},
        ],
        "sizes": ["CH", "M", "G", "XG"],
    },
    "CANADA": {
        "aliases": ["CANADA", "CAN", "COPO26CAN"],
        "colors": [
            {"name": "NEGRO", "code": "01"},
            {"name": "ROJO", "code": "02"},
            {"name": "BLANCO", "code": "03"},
        ],
        "sizes": ["CH", "M", "G", "XG"],
    },
    "MONACO": {
        "aliases": ["MONACO", "MON", "COPO26MON"],
        "colors": [
            {"name": "BLANCO", "code": "03"},
            {"name": "MARINO", "code": "04"},
            {"name": "NEGRO", "code": "01"},
        ],
        "sizes": ["CH", "M", "G", "XG"],
    },
    "MEXICO": {
        "aliases": ["MEXICO", "MEX", "COPO26MEX", "CC-MEX"],
        "colors": [
            {"name": "NEGRO", "code": "01"},
            {"name": "BLANCO", "code": "03"},
            {"name": "VERDE", "code": "17"},
        ],
        "sizes": ["CH", "M", "G", "XG"],
    },
    "NETHERLANDS": {
        "aliases": ["NETHERLANDS", "NED", "HOLANDA", "CC-NED"],
        "colors": [
            {"name": "NARANJA", "code": "09"},
            {"name": "NEGRO", "code": "01"},
        ],
        "sizes": ["CH", "M", "G", "XG"],
    },
    "24H LE MANS": {
        "aliases": ["24H", "LE MANS", "24H-101"],
        "colors": [
            {"name": "NEGRO", "code": "01"},
            {"name": "BLANCO", "code": "03"},
            {"name": "MARINO", "code": "04"},
            {"name": "ROJO", "code": "02"},
        ],
        "sizes": ["CH", "M", "G", "XG"],
    },
    "MX RACING 102": {
        "aliases": ["MXR-24-102", "102"],
        "colors": [
            {"name": "NEGRO", "code": "01"},
            {"name": "ROJO", "code": "02"},
            {"name": "BLANCO", "code": "03"},
            {"name": "MARINO", "code": "04"},
        ],
        "sizes": ["CH", "M", "G", "XG"],
    },
    "GORRA ALFA ROMEO": {
        "aliases": ["AR-CAP", "26910", "ALFA ROMEO"],
        "colors": [
            {"name": "NEGRO", "code": "01"},
            {"name": "ROJO", "code": "02"},
            {"name": "BLANCO", "code": "03"},
            {"name": "MARINO", "code": "04"},
            {"name": "GRIS", "code": "07"},
            {"name": "GRIS OSCURO", "code": "07O"},
        ],
        "sizes": ["Unitalla"],
    },
}


def load_catalog(filepath: Optional[Path] = None) -> dict:
    target = filepath or CATALOG_PATH
    if not target.exists():
        save_catalog(DEFAULT_CATALOG, target)
        return json.loads(json.dumps(DEFAULT_CATALOG))
    try:
        with open(target, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return json.loads(json.dumps(DEFAULT_CATALOG))


def save_catalog(catalog: dict, filepath: Optional[Path] = None) -> None:
    target = filepath or CATALOG_PATH
    target.parent.mkdir(parents=True, exist_ok=True)
    with open(target, "w", encoding="utf-8") as f:
        json.dump(catalog, f, indent=4, ensure_ascii=False)


def normalize_size(raw_size: str) -> str:
    cleaned = (raw_size or "").strip().upper()
    return SIZE_NORMALIZATION.get(cleaned, cleaned or "U")


def parse_item_info(item: OrderItem, catalog: dict) -> Tuple[str, str, str, str]:
    """
    Determina (model_name, color_name, color_code, size) a partir de OrderItem y el catálogo.
    """
    title_upper = (item.title or "").upper()
    sku_upper = (item.sku_oferta or "").upper()
    sku_parts = [p.strip() for p in re.split(r"[-_]+", sku_upper) if p.strip()]

    # 1. Identificar modelo buscando en catálogo
    matched_model_name: Optional[str] = None
    for model_name, m_info in catalog.items():
        aliases = [model_name.upper()] + [a.upper() for a in m_info.get("aliases", [])]
        for alias in aliases:
            # Match por palabra o subcadena en SKU o título
            pattern = rf"\b{re.escape(alias)}\b"
            if re.search(pattern, sku_upper) or re.search(pattern, title_upper):
                matched_model_name = model_name
                break
            if alias in sku_upper:
                matched_model_name = model_name
                break
        if matched_model_name:
            break

    # Si no hubo match, deducir nombre de modelo del SKU o título
    if not matched_model_name:
        if len(sku_parts) >= 3:
            matched_model_name = f"MODELO {sku_parts[1]}"
        elif item.title:
            matched_model_name = " ".join(item.title.split()[:4]).upper()
        else:
            matched_model_name = "MODELO GENERAL"

    model_entry = catalog.get(matched_model_name, {})
    model_colors = model_entry.get("colors", [])
    code_to_name = {c["code"].upper(): c["name"].upper() for c in model_colors}

    # 2. Extraer Talla
    size = ""
    if item.talla:
        size = normalize_size(item.talla)
    elif sku_parts:
        last_part = sku_parts[-1]
        if last_part in SIZE_NORMALIZATION or last_part in ["CH", "M", "G", "XG", "S", "L", "XL", "2XL", "XXL"]:
            size = normalize_size(last_part)
    if not size:
        size = "CH"

    # 3. Extraer Código de Color y Nombre de Color
    color_code = ""
    color_name = ""

    # Buscar código en los fragmentos del SKU (ej: 01, 04, 07O, 15)
    candidates_in_sku = []
    for part in sku_parts:
        if re.fullmatch(r"\d{2}[A-Za-z]?", part):
            candidates_in_sku.append(part)

    if candidates_in_sku:
        color_code = candidates_in_sku[-1]

    # Resolver color_name a partir de color_code
    if color_code:
        if color_code in code_to_name:
            color_name = code_to_name[color_code]
        elif color_code in DEFAULT_COLORS:
            color_name = DEFAULT_COLORS[color_code]

    # Si no se encontró por código, buscar por nombre en el título o SKU
    if not color_name:
        for name, code in [(c["name"], c["code"]) for c in model_colors] + list(DEFAULT_COLORS.items()):
            if isinstance(code, str) and code.isalpha():
                name, code = code, name
            if re.search(rf"\b{re.escape(name)}\b", title_upper) or re.search(rf"\b{re.escape(name)}\b", sku_upper):
                color_name = name
                color_code = code
                break

    if not color_code:
        color_code = "01"
    if not color_name:
        color_name = DEFAULT_COLORS.get(color_code, f"COLOR {color_code}")

    return matched_model_name, color_name, color_code, size


def generate_models_sales_excel(
    batch: DayBatch,
    output_path: Path,
    catalog: Optional[dict] = None,
) -> Path:
    """
    Genera el archivo Excel de ventas desglosadas por modelo con el formato:
    [MODELO]
    COLOR | CÓDIGO COLOR | CH | M | G | XG | TOTAL
    NEGRO | 01           | 0  | 0 | 0 | 0  | 0
    ...
    """
    if catalog is None:
        catalog = load_catalog()

    ok_orders = [o for o in batch.orders if o.status == "ok"]

    # Estructura: sales[model_name][(color_name, color_code)][size] = qty
    sales: Dict[str, Dict[Tuple[str, str], Dict[str, int]]] = {}

    for order in ok_orders:
        for item in order.items:
            m_name, c_name, c_code, size = parse_item_info(item, catalog)
            if m_name not in sales:
                sales[m_name] = {}
            key = (c_name, c_code)
            if key not in sales[m_name]:
                sales[m_name][key] = {}
            sales[m_name][key][size] = sales[m_name][key].get(size, 0) + item.qty

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ventas por Modelo"
    ws.views.sheetView[0].showGridLines = True

    # Paleta de estilos
    title_font = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
    title_fill = PatternFill(start_color="1B4F72", end_color="1B4F72", fill_type="solid")

    header_font = Font(name="Segoe UI", size=11, bold=True, color="1A252C")
    header_fill = PatternFill(start_color="D6EAF8", end_color="D6EAF8", fill_type="solid")

    bold_font = Font(name="Segoe UI", size=11, bold=True)
    normal_font = Font(name="Segoe UI", size=11)
    muted_zero_font = Font(name="Segoe UI", size=11, color="7F8C8D")

    total_fill = PatternFill(start_color="EAECEE", end_color="EAECEE", fill_type="solid")

    thin_border_side = Side(style="thin", color="BDC3C7")
    cell_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side,
    )

    double_bottom_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=Side(style="double", color="2C3E50"),
    )

    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    current_row = 1

    # Título general
    ws.cell(row=current_row, column=1, value=f"RESUMEN DE VENTAS POR MODELO - FECHA: {batch.date}")
    ws.cell(row=current_row, column=1).font = Font(name="Segoe UI", size=15, bold=True, color="1B4F72")
    current_row += 2

    # Modelos a mostrar:
    # 1. Modelos registrados en catálogo que tengan ventas
    # 2. Modelos no registrados que tuvieron ventas
    # 3. Si no hubo ventas, mostrar los modelos registrados principales en ceros
    models_to_display = []
    if sales:
        # Priorizar modelos en orden del catálogo, luego extras
        for m in catalog.keys():
            if m in sales:
                models_to_display.append(m)
        for m in sales.keys():
            if m not in models_to_display:
                models_to_display.append(m)
    else:
        # Si no hay ventas 'ok', mostramos modelos del catálogo como plantilla
        models_to_display = list(catalog.keys())[:3]

    grand_total_pieces = 0

    for m_name in models_to_display:
        m_catalog = catalog.get(m_name, {})
        m_sales = sales.get(m_name, {})

        # Tallas de la tabla (usar las del catálogo o por defecto)
        default_sizes = m_catalog.get("sizes", ["CH", "M", "G", "XG"])
        # Añadir tallas vendidas extras si existieran
        used_sizes = list(default_sizes)
        for c_dict in m_sales.values():
            for s in c_dict.keys():
                if s not in used_sizes:
                    used_sizes.append(s)

        num_size_cols = len(used_sizes)
        total_cols = 2 + num_size_cols + 1  # Color, Código, tallas..., Total

        # 1. Fila de Título de Modelo (Banner)
        ws.merge_cells(
            start_row=current_row,
            start_column=1,
            end_row=current_row,
            end_column=total_cols,
        )
        banner_cell = ws.cell(row=current_row, column=1, value=m_name.upper())
        banner_cell.font = title_font
        banner_cell.fill = title_fill
        banner_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[current_row].height = 26

        for col in range(1, total_cols + 1):
            ws.cell(row=current_row, column=col).border = cell_border
        current_row += 1

        # 2. Fila de Encabezados de Columna
        headers = ["COLOR", "CÓDIGO COLOR"] + used_sizes + ["TOTAL"]
        ws.row_dimensions[current_row].height = 22
        for col_idx, h_text in enumerate(headers, start=1):
            c = ws.cell(row=current_row, column=col_idx, value=h_text)
            c.font = header_font
            c.fill = header_fill
            c.alignment = center_align if col_idx > 1 else left_align
            c.border = cell_border
        current_row += 1

        # 3. Filas de Colores
        # Comenzamos con los colores registrados en catálogo para este modelo
        color_rows_order: List[Tuple[str, str]] = []
        registered_colors = m_catalog.get("colors", [])
        for rc in registered_colors:
            color_rows_order.append((rc["name"].upper(), str(rc["code"]).strip()))

        # Añadir colores vendidos que no estuvieran en catálogo
        for (c_name, c_code) in m_sales.keys():
            if (c_name.upper(), c_code) not in color_rows_order:
                color_rows_order.append((c_name.upper(), c_code))

        col_totals = [0] * num_size_cols
        model_grand_total = 0

        for (c_name, c_code) in color_rows_order:
            ws.row_dimensions[current_row].height = 20
            # Nombre de color
            c1 = ws.cell(row=current_row, column=1, value=c_name)
            c1.font = normal_font
            c1.border = cell_border
            c1.alignment = left_align

            # Código de color
            c2 = ws.cell(row=current_row, column=2, value=str(c_code).zfill(2) if c_code.isdigit() and len(c_code) == 1 else str(c_code))
            c2.font = normal_font
            c2.border = cell_border
            c2.alignment = center_align

            row_qty_sum = 0
            # Columnas de tallas
            for s_idx, size_col in enumerate(used_sizes):
                q = m_sales.get((c_name, c_code), {}).get(size_col, 0)
                row_qty_sum += q
                col_totals[s_idx] += q

                cq = ws.cell(row=current_row, column=3 + s_idx, value=q)
                cq.font = normal_font if q > 0 else muted_zero_font
                cq.alignment = center_align
                cq.border = cell_border

            # Total de la fila
            ct = ws.cell(row=current_row, column=total_cols, value=row_qty_sum)
            ct.font = bold_font if row_qty_sum > 0 else muted_zero_font
            ct.alignment = center_align
            ct.border = cell_border

            model_grand_total += row_qty_sum
            current_row += 1

        # 4. Fila de Totales del Modelo
        ws.row_dimensions[current_row].height = 22
        tot_lbl = ws.cell(row=current_row, column=1, value="TOTAL")
        tot_lbl.font = bold_font
        tot_lbl.fill = total_fill
        tot_lbl.alignment = left_align
        tot_lbl.border = double_bottom_border

        tot_code = ws.cell(row=current_row, column=2, value="")
        tot_code.fill = total_fill
        tot_code.border = double_bottom_border

        for s_idx, size_tot in enumerate(col_totals):
            cst = ws.cell(row=current_row, column=3 + s_idx, value=size_tot)
            cst.font = bold_font
            cst.fill = total_fill
            cst.alignment = center_align
            cst.border = double_bottom_border

        cmt = ws.cell(row=current_row, column=total_cols, value=model_grand_total)
        cmt.font = bold_font
        cmt.fill = total_fill
        cmt.alignment = center_align
        cmt.border = double_bottom_border

        grand_total_pieces += model_grand_total
        current_row += 3  # Espacio entre modelos

    # Resumen final de piezas
    ws.merge_cells(
        start_row=current_row,
        start_column=1,
        end_row=current_row,
        end_column=4,
    )
    gt_cell = ws.cell(row=current_row, column=1, value=f"TOTAL PIEZAS VENDIDAS EN EL DÍA: {grand_total_pieces}")
    gt_cell.font = Font(name="Segoe UI", size=12, bold=True, color="1B4F72")
    gt_cell.alignment = left_align

    # Ajustar ancho de columnas automáticamente
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if cell.coordinate in ws.merged_cells:
                continue
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 16

    wb.save(output_path)
    return output_path
